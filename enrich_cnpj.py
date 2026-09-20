"""
CNPJ Data Enrichment

Enriches a spreadsheet of Brazilian company registrations (CNPJ) with public
records from the Receita Federal, via the BrasilAPI public endpoint.

Designed for vendor due diligence: it fills in registered name, trade name,
registration status, primary economic activity (CNAE), business segment,
company size and public contact details, and raises a risk flag when a
registration is not active.

The process is resumable. Only rows marked PENDING are processed, and the
workbook is saved periodically, so an interrupted run can be restarted
without duplicating requests.

Usage:
    pip install openpyxl requests
    python enrich_cnpj.py vendors.xlsx --sheet VENDORS --segments segments.json

Author: Daniel Batista
License: MIT
"""

import argparse
import json
import re
import sys
import time
import unicodedata
from datetime import date
from pathlib import Path

import requests
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

API_URL = "https://brasilapi.com.br/api/cnpj/v1/{cnpj}"
USER_AGENT = "cnpj-enricher/1.0 (+https://github.com/yl-dan/cnpj-enricher)"

REQUEST_TIMEOUT = 20       # seconds
THROTTLE_SECONDS = 1.2     # delay between requests, to stay under rate limits
BACKOFF_SECONDS = 30       # wait after HTTP 429
MAX_ATTEMPTS = 3           # attempts per row before giving up
SAVE_EVERY = 25            # rows between checkpoint saves

STATUS_PENDING = "PENDING"
STATUS_DONE = "ENRICHED"
STATUS_NOT_FOUND = "NOT FOUND"
STATUS_FAILED = "REQUEST FAILED"

# Columns the sheet must contain. Missing columns are reported before any
# request is made, rather than raising a KeyError halfway through a long run.
REQUIRED_COLUMNS = [
    "cnpj",
    "status_enrichment",
    "registered_name",
    "trade_name",
    "registration_status",
    "primary_cnae",
    "segment",
    "incorporation_date",
    "company_size",
    "share_capital",
    "city",
    "state",
    "public_email",
    "public_phone",
    "risk_flag",
    "source",
    "query_date",
]

# Fallback segment map. Keys are CNAE prefixes, matched longest-first.
# Replace with your own via --segments; see segments.example.json.
DEFAULT_SEGMENTS = {
    "10.": "Food manufacturing",
    "17.": "Paper and packaging manufacturing",
    "46.": "Wholesale trade",
    "47.": "Retail trade",
    "49.": "Transport",
    "52.": "Logistics and warehousing",
    "62.": "Information technology",
    "69.": "Legal and accounting services",
}
SEGMENT_FALLBACK = "Other"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def strip_accents(text):
    """
    Removes diacritics so comparisons are accent-insensitive.

    Registry data is inconsistent about accents: the same status can arrive as
    "RECUPERACAO JUDICIAL" or "RECUPERAÇÃO JUDICIAL". Comparing raw strings
    silently misses the second form, which is exactly the case that matters.
    """
    normalized = unicodedata.normalize("NFKD", text or "")
    return "".join(c for c in normalized if not unicodedata.combining(c))


def digits_only(value):
    return re.sub(r"\D", "", str(value or ""))


def format_cnae(raw_code):
    """Formats a 7-digit CNAE code as NN.NN-N-NN. Returns '' if absent."""
    code = str(raw_code or "").zfill(7)
    if not code.strip("0"):
        return ""
    return f"{code[:2]}.{code[2:4]}-{code[4]}-{code[5:]}"


def format_phone(raw_phone):
    """Formats a Brazilian phone number. Returns '' if too short to be valid."""
    phone = digits_only(raw_phone)
    if len(phone) < 10:
        return ""
    return f"({phone[:2]}) {phone[2:-4]}-{phone[-4:]}"


def format_date(iso_date):
    """Converts YYYY-MM-DD to DD/MM/YYYY. Returns '' on anything unexpected."""
    try:
        year, month, day = str(iso_date).split("-")
        return f"{day}/{month}/{year}"
    except (ValueError, AttributeError):
        return ""


def classify_segment(cnae, segments):
    """Maps a formatted CNAE to a business segment by longest prefix match."""
    if not cnae:
        return ""
    for prefix in sorted(segments, key=len, reverse=True):
        if cnae.startswith(prefix):
            return segments[prefix]
    return SEGMENT_FALLBACK


def load_segments(path):
    if not path:
        return DEFAULT_SEGMENTS
    try:
        with open(path, encoding="utf-8") as handle:
            mapping = json.load(handle)
    except (OSError, json.JSONDecodeError) as error:
        sys.exit(f"Could not read segment map '{path}': {error}")
    if not isinstance(mapping, dict):
        sys.exit(f"Segment map '{path}' must be a JSON object of prefix -> label.")
    return mapping


# ---------------------------------------------------------------------------
# API access
# ---------------------------------------------------------------------------

def fetch_company(session, cnpj):
    """
    Retrieves one company record.

    Returns (status, payload):
        ("ok", dict)        record found
        ("not_found", None) the registry has no such CNPJ
        ("failed", None)    exhausted retries on network errors or throttling

    Rate limiting and transient failures are retried against the same CNPJ
    rather than skipping to the next row, so one slow moment does not leave
    gaps scattered through the output.
    """
    for attempt in range(1, MAX_ATTEMPTS + 1):
        try:
            response = session.get(
                API_URL.format(cnpj=cnpj),
                timeout=REQUEST_TIMEOUT,
            )
        except requests.RequestException:
            time.sleep(BACKOFF_SECONDS if attempt > 1 else 5)
            continue

        if response.status_code == 200:
            try:
                return "ok", response.json()
            except ValueError:
                return "failed", None

        if response.status_code == 404:
            return "not_found", None

        if response.status_code == 429 or response.status_code >= 500:
            time.sleep(BACKOFF_SECONDS * attempt)
            continue

        return "failed", None

    return "failed", None


def build_updates(payload, segments, today):
    """Maps an API payload onto the sheet's column names."""
    cnae = format_cnae(payload.get("cnae_fiscal"))
    cnae_label = payload.get("cnae_fiscal_descricao") or ""

    registration_status = payload.get("descricao_situacao_cadastral") or ""
    registered_name = payload.get("razao_social") or ""

    flags = []
    normalized_status = strip_accents(registration_status).upper()
    if normalized_status and normalized_status != "ATIVA":
        flags.append(f"REGISTRATION {normalized_status}")
    if "RECUPERACAO JUDICIAL" in strip_accents(registered_name).upper():
        flags.append("JUDICIAL RECOVERY")

    return {
        "registered_name": registered_name,
        "trade_name": payload.get("nome_fantasia"),
        "registration_status": registration_status,
        "primary_cnae": f"{cnae} - {cnae_label}" if cnae else "",
        "segment": classify_segment(cnae, segments),
        "incorporation_date": format_date(payload.get("data_inicio_atividade")),
        "company_size": payload.get("porte"),
        "share_capital": payload.get("capital_social"),
        "city": payload.get("municipio"),
        "state": payload.get("uf"),
        "public_email": (payload.get("email") or "").lower(),
        "public_phone": format_phone(payload.get("ddd_telefone_1")),
        "risk_flag": "; ".join(flags),
        "status_enrichment": STATUS_DONE,
        "source": "BrasilAPI (Receita Federal)",
        "query_date": today,
    }


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def parse_args():
    parser = argparse.ArgumentParser(
        description="Enrich a CNPJ spreadsheet with public Receita Federal data.",
    )
    parser.add_argument("workbook", help="path to the .xlsx file to enrich")
    parser.add_argument("--sheet", default="VENDORS", help="worksheet name (default: VENDORS)")
    parser.add_argument("--segments", help="path to a JSON map of CNAE prefix -> segment label")
    parser.add_argument("--limit", type=int, help="stop after this many rows (useful for testing)")
    return parser.parse_args()


def main():
    args = parse_args()

    path = Path(args.workbook)
    if not path.is_file():
        sys.exit(f"File not found: {path}")

    segments = load_segments(args.segments)

    workbook = load_workbook(path)
    if args.sheet not in workbook.sheetnames:
        sys.exit(f"Sheet '{args.sheet}' not found. Available: {', '.join(workbook.sheetnames)}")
    sheet = workbook[args.sheet]

    columns = {cell.value: cell.column for cell in sheet[1] if cell.value}
    missing = [name for name in REQUIRED_COLUMNS if name not in columns]
    if missing:
        sys.exit("Missing required column(s): " + ", ".join(missing))

    pending = [
        row for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, columns["status_enrichment"]).value == STATUS_PENDING
    ]
    if args.limit:
        pending = pending[:args.limit]

    if not pending:
        print("Nothing to do: no rows marked PENDING.")
        return

    print(f"{len(pending)} row(s) to enrich.")

    today = date.today().isoformat()
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})

    counts = {"enriched": 0, "not_found": 0, "failed": 0, "invalid": 0}

    for processed, row in enumerate(pending, start=1):
        cnpj = digits_only(sheet.cell(row, columns["cnpj"]).value)

        if len(cnpj) != 14:
            sheet.cell(row, columns["status_enrichment"]).value = "INVALID CNPJ"
            sheet.cell(row, columns["query_date"]).value = today
            counts["invalid"] += 1
        else:
            status, payload = fetch_company(session, cnpj)

            if status == "ok":
                for name, value in build_updates(payload, segments, today).items():
                    if value not in (None, ""):
                        sheet.cell(row, columns[name]).value = value
                # status and date are always written, even when blank above
                sheet.cell(row, columns["status_enrichment"]).value = STATUS_DONE
                counts["enriched"] += 1
            elif status == "not_found":
                sheet.cell(row, columns["status_enrichment"]).value = STATUS_NOT_FOUND
                counts["not_found"] += 1
            else:
                # Left as FAILED rather than PENDING so the run is auditable;
                # set it back to PENDING to retry these rows on a later run.
                sheet.cell(row, columns["status_enrichment"]).value = STATUS_FAILED
                counts["failed"] += 1

            sheet.cell(row, columns["query_date"]).value = today
            time.sleep(THROTTLE_SECONDS)

        if processed % SAVE_EVERY == 0:
            workbook.save(path)
            print(f"  {processed}/{len(pending)} processed, checkpoint saved")

    workbook.save(path)
    print(
        "Done. "
        f"enriched={counts['enriched']} "
        f"not_found={counts['not_found']} "
        f"failed={counts['failed']} "
        f"invalid={counts['invalid']}"
    )


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit("\nInterrupted. Progress up to the last checkpoint was saved.")
